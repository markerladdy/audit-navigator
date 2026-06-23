from pptx import Presentation
import json

p = r"C:\Users\MuMa Studio\auditnavigator\downloads\用友高校实训-智能制造.pptx"
prs = Presentation(p)

slides = []
for slide in prs.slides:
    items = []
    for shape in slide.shapes:
        entry = {}
        if shape.has_text_frame:
            paras = []
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    paras.append({"t": t, "l": para.level if para.level else 0})
            if paras:
                entry["p"] = paras
        if shape.has_table:
            rows = []
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            if rows:
                entry["t"] = rows
        if entry:
            items.append(entry)
    slides.append(items)

with open(r"C:\Users\MuMa Studio\auditnavigator\downloads\slides_data.json", "w", encoding="utf-8") as f:
    json.dump({"total": len(slides), "slides": slides}, f, ensure_ascii=False)

print(f"Done: {len(slides)} slides exported")
