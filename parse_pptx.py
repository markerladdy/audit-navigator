from pptx import Presentation
import json, re

pptx_path = "/c/Users/MuMa Studio/auditnavigator/downloads/用友高校实训-智能制造.pptx"
prs = Presentation(pptx_path)

slides_data = []
for i, slide in enumerate(prs.slides):
    slide_info = {"index": i + 1, "elements": []}
    for shape in slide.shapes:
        item = {"type": str(shape.shape_type), "left": shape.left, "top": shape.top, "width": shape.width, "height": shape.height}
        if shape.has_text_frame:
            paras = []
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    level = para.level if para.level else 0
                    paras.append({"text": text, "level": level})
            if paras:
                item["paragraphs"] = paras
        if shape.has_table:
            table_data = []
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                table_data.append(cells)
            item["table"] = table_data
        slide_info["elements"].append(item)
    slides_data.append(slide_info)

print(f"Total slides: {len(slides_data)}")
print(f"Total elements: {sum(len(s['elements']) for s in slides_data)}")

# Print slide titles for overview
for s in slides_data:
    title = ""
    for e in s["elements"]:
        if e.get("paragraphs"):
            for p in e["paragraphs"]:
                if p["level"] == 0 and len(p["text"]) > 3:
                    title = p["text"][:60]
                    break
        if title:
            break
    print(f"  Slide {s['index']}: {title}")
